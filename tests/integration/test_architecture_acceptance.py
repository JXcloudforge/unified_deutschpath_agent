from pathlib import Path

import pytest
from docx import Document

from app.agent import UnifiedOrchestrator
from app.chat import ActorContext, ChatService
from app.chat.repository import ChatRepository
from app.files import FileService
from app.llm import FakeLLM
from app.rag import DeterministicEmbedding, RAGService
from app.rag.vector_store import ChromaVectorStore
from app.schemas import ParsedDocument
from app.users import UserStateRepository, UserStateService


def build_stack(tmp_path: Path, llm=None):
    rag = RAGService(ChromaVectorStore(tmp_path / "chroma", DeterministicEmbedding()))
    users = UserStateService(UserStateRepository(tmp_path / "state.db"))
    return rag, users, llm or FakeLLM("FAKE_ANSWER")


def test_docx_direct_prompt_e2e(tmp_path: Path):
    path = tmp_path / "secret.docx"
    word = Document(); word.add_paragraph("STAGE2_DOCX_SECRET_88421"); word.save(path)
    rag, users, llm = build_stack(tmp_path)
    document = FileService().parse(path)
    UnifiedOrchestrator(rag, users, llm).answer(user_id="u", message="Summarize this document", attachment=document)
    assert "STAGE2_DOCX_SECRET_88421" in llm.last_prompt


def test_session_and_knowledge_namespaces_are_strictly_isolated(tmp_path: Path):
    rag, _, _ = build_stack(tmp_path)
    rag.index_documents("session:A", [ParsedDocument("a.txt", "txt", "ALPHA_SESSION_SECRET_123", {"source": "A"})])
    rag.index_documents("session:B", [ParsedDocument("b.txt", "txt", "BETA_SESSION_SECRET_456", {"source": "B"})])
    rag.index_documents("knowledge:german", [ParsedDocument("german.txt", "txt", "GERMAN_SECRET_111", {"source": "german"})])
    rag.index_documents("knowledge:germany", [ParsedDocument("germany.txt", "txt", "GERMANY_SECRET_222", {"source": "germany"})])
    a = rag.search("session:A", "ALPHA_SESSION_SECRET_123", 3)
    b = rag.search("session:B", "BETA_SESSION_SECRET_456", 3)
    german = rag.search("knowledge:german", "GERMAN_SECRET_111", 3)
    germany = rag.search("knowledge:germany", "GERMANY_SECRET_222", 3)
    assert all("BETA" not in item.text for item in a) and all("ALPHA" not in item.text for item in b)
    assert all("GERMANY" not in item.text for item in german) and all("GERMAN_SECRET_111" not in item.text for item in germany)
    assert german[0].metadata["namespace"] == "knowledge:german"
    assert germany[0].metadata["filename"] == "germany.txt"


def test_chat_upload_to_prompt_to_persistence_e2e(tmp_path: Path):
    path = tmp_path / "chat.docx"
    word = Document(); word.add_paragraph("CHAT_DOCX_SECRET_999"); word.save(path)
    rag, users, llm = build_stack(tmp_path)
    repository = ChatRepository(tmp_path / "chat.db")
    chat = ChatService(repository, UnifiedOrchestrator(rag, users, llm), FileService())
    actor = ActorContext.from_local_header("alice")
    session = chat.create_session(actor)
    chat.upload(actor, session.id, path)
    result = chat.send_message(actor, session.id, "What does the file say?")
    messages = chat.get_messages(actor, session.id)
    assert result.assistant_message.content == "FAKE_ANSWER"
    assert [message.role for message in messages] == ["user", "assistant"]
    assert "CHAT_DOCX_SECRET_999" in llm.last_prompt
    assert repository.attachments("alice", session.id)[0]["session_id"] == session.id
    with pytest.raises(KeyError): chat.get_messages(ActorContext.from_local_header("bob"), session.id)


def test_llm_failure_keeps_user_message_without_fake_assistant(tmp_path: Path):
    class ExplodingLLM:
        def answer(self, prompt: str) -> str:
            raise RuntimeError("provider unavailable")
    rag, users, _ = build_stack(tmp_path)
    repository = ChatRepository(tmp_path / "chat.db")
    chat = ChatService(repository, UnifiedOrchestrator(rag, users, ExplodingLLM()), FileService())
    actor = ActorContext.from_local_header("alice")
    session = chat.create_session(actor)
    with pytest.raises(RuntimeError, match="agent service failed"):
        chat.send_message(actor, session.id, "Hallo")
    assert [message.role for message in chat.get_messages(actor, session.id)] == ["user"]
