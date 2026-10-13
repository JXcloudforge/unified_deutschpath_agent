from pathlib import Path

from docx import Document

from app.schemas import ParsedDocument


class DocxParser:
    extensions = (".docx",)

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return ParsedDocument(path.name, "docx", "\n".join(paragraphs), {"source": str(path)})
